from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Pt
import copy

class SlideBlocksEngine:
    """Slide Blocks 核心引擎：处理布局计算、样式同步与元素克隆。"""

    @staticmethod
    def calculate_grid_layout(total_width, columns, gap=200000, margin_left=450000):
        """计算 N 分类布局的精确 X 轴坐标。"""
        # 计算每个 Block 的可用宽度
        usable_width = total_width - 2 * margin_left - (columns - 1) * gap
        block_width = int(usable_width / columns)
        
        # 生成每个列的起始位置 (Left)
        lefts = [int(margin_left + i * (block_width + gap)) for i in range(columns)]
        return lefts, block_width

    @staticmethod
    def sync_font_style(source_tf, target_tf, is_title=False):
        """深度同步字体基因：字号、字体、加粗、斜体、颜色。"""
        target_tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        target_tf.word_wrap = True
        target_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        try:
            source_run = source_tf.paragraphs[0].runs[0]
            s_font = source_run.font
            s_name, s_size, s_bold, s_italic = s_font.name, s_font.size, s_font.bold, s_font.italic
            s_rgb = s_font.color.rgb if hasattr(s_font.color, 'rgb') else None
        except:
            # 降级方案
            s_name, s_size, s_bold, s_italic, s_rgb = '微软雅黑', Pt(18) if is_title else Pt(11), is_title, False, None

        for paragraph in target_tf.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.name = s_name
                run.font.size = s_size
                run.font.bold = s_bold
                run.font.italic = s_italic
                if s_rgb:
                    run.font.color.rgb = s_rgb

    @staticmethod
    def deep_clone_shape(shape, slide, new_left):
        """利用底层 XML 深度克隆形状，并设置新坐标。"""
        el = shape.element
        new_el = copy.deepcopy(el)
        slide.shapes._spTree.append(new_el)
        cloned_shape = slide.shapes[-1]
        cloned_shape.left = int(new_left)
        return cloned_shape
